import random
import openpyxl

affiliates = {
	"france": {"min": 0, "max":20000, "products": ["P1", "P2", "P3", "P4"], "code": "001" },
	"spain": {"min": 0, "max":10000, "products": ["P1", "P2"], "code": "002" },
	"chili": {"min": 0, "max":5000, "products": ["P1", "P3"] , "code": "003"},
	"australia": {"min": 0, "max":20000, "products": ["P1", "P2"], "code": "004" }
}

affiliate_rand_forcast_model = openpyxl.load_workbook('sales_forcast_affiliate_model.xlsx')

def accumu(lis):
    total = 0
    for v in lis:
        total += v
        yield total

def randSalesForcast(a, horizon):
    min = affiliates[a]['min']
    max = affiliates[a]['max']
    ans = [random.randint(min, max) for _ in range(horizon)]
    return ans

def getModelValues(affiliate, t):
    model_sheet = affiliate_rand_forcast_model.get_sheet_by_name(affiliates[affiliate]["code"])
    a = int(model_sheet.cell(row=2, column=t + 2).value)
    b = int(model_sheet.cell(row=3, column=t + 2).value)   
    c = int(model_sheet.cell(row=4, column=t + 2).value) 
    d = int(model_sheet.cell(row=5, column=t + 2).value)
    ref_week = int(model_sheet.cell(row=6, column=t + 2).value)
    return a, b, c, d, ref_week

def pickRand(min1, min2, max1, max2, prob):
    if random.random() < prob:
        return random.randint(min2, max2)
    else:
        if random.random() < 0.5:
            return random.randint(min1, min2)
        else:
            return random.randint(max2, max1)
        
def calcRandCPV(horizon, affiliate, prev_pv):
    prev_cpv = list(accumu(prev_pv))
    acpv = [0 for _ in range(horizon)]
    for t in range(horizon):
        min_pv = affiliates[affiliate]['min']
        max_pv = affiliates[affiliate]['max']
        cpv = prev_cpv[t+1] if t + 1 < horizon else prev_cpv[horizon-1] + random.randint(min_pv, max_pv)
        a, b, c, d, rw = getModelValues(affiliate, t)
        min_1 = cpv + a * (cpv - prev_cpv[rw-1])
        min_2 = cpv + b * (cpv - prev_cpv[rw-1]) 
        max_2 = cpv + c * (cpv - prev_cpv[rw-1])
        max_1 = cpv + d * (cpv - prev_cpv[rw-1]) 
        rd = pickRand(min_1, min_2, max_1, max_2, 0.8)
        acpv[t] =  max(rd, acpv[t-1] if t>0 else 0)
    return acpv
        
def genRandSalesForcast(horizon, affiliate, prev_pv=None):
    if prev_pv is None:
        return randSalesForcast(affiliate, horizon)
    acpv = calcRandCPV(horizon, affiliate, prev_pv)
    new_apv = [None for _ in range(horizon)]
    new_apv[0] = acpv[0] - prev_pv[0]
    for t in range(1, horizon):
        new_apv[t] = acpv[t] - acpv[t-1]
    return new_apv

def run(prev_sales_forcast, horizon):
    return {
        a: {
            p: genRandSalesForcast(horizon, a, prev_sales_forcast[a][p]) for p in aff["products"]
        } for a, aff in affiliates.items()
    }
